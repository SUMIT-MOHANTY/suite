from rest_framework import serializers
from openpyxl import load_workbook
from django.http import HttpResponse
import io, hashlib, os

from .models import Policy
from .constants import PolicyStatus

class PolicySerializer(serializers.ModelSerializer):
    class Meta:
        model = Policy
        fields = "__all__"
        read_only_fields = ("created_at", "updated_at")

class ACORDExcelSerializer(serializers.ModelSerializer):
    class Meta:
        model  = Policy
        fields = ("id", "policy_number", "effective_date", "expiry_date", "premium")

    def generate_xlsx(self, instance: Policy):
        template_path = os.path.join(os.path.dirname(__file__), "openpyxl-templates", "sample_acord_125.xlsx")
        wb = load_workbook(template_path)
        ws = wb.active
        # Fill stub placeholders
        ws["B4"] = instance.policy_number
        ws["D4"] = instance.effective_date.strftime("%Y-%m-%d")
        ws["F4"] = instance.expiry_date.strftime("%Y-%m-%d")
        ws["H4"] = float(instance.premium)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf
